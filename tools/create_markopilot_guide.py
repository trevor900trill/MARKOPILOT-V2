from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path("docs/Markopilot-How-It-Works-Guide.docx")
PURPLE = "7C6EFF"; INK = "17151F"; MUTED = "6B6877"; PALE = "F1EFFF"; GREEN = "167957"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tcPr.append(shd)
def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); mar = tcPr.first_child_found_in("w:tcMar")
    if mar is None: mar = OxmlElement("w:tcMar"); tcPr.append(mar)
    for side, val in (("top",top),("start",start),("bottom",bottom),("end",end)):
        node = mar.find(qn(f"w:{side}"))
        if node is None: node = OxmlElement(f"w:{side}"); mar.append(node)
        node.set(qn("w:w"), str(val)); node.set(qn("w:type"), "dxa")
def font(run, size=11, color=INK, bold=False, italic=False):
    run.font.name = "Calibri"; run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri"); run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size); run.font.color.rgb = RGBColor.from_string(color); run.bold = bold; run.italic = italic
def add_text(p, text, **kw):
    r = p.add_run(text); font(r, **kw); return r
def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}"); p.paragraph_format.keep_with_next = True; add_text(p, text, size={1:16,2:13,3:12}[level], color=PURPLE if level<3 else "3D356E", bold=True); return p
def body(doc, text):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6); p.paragraph_format.line_spacing = 1.1; add_text(p,text); return p
def bullet(doc, text):
    p=doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after=Pt(4); p.paragraph_format.line_spacing=1.1; add_text(p,text); return p
def callout(doc, title, text):
    t=doc.add_table(rows=1, cols=1); t.autofit=False; t.columns[0].width=Inches(6.5); c=t.cell(0,0); shade(c,PALE); set_cell_margins(c,160,180,160,180); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(3); add_text(p,title.upper(),size=9,color=PURPLE,bold=True)
    p=c.add_paragraph(); p.paragraph_format.space_after=Pt(0); add_text(p,text,size=10.5,color="322E45")
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
def numbered(doc, number, title, text):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(2); add_text(p,f"{number}. ",size=12,color=PURPLE,bold=True); add_text(p,title,size=12,bold=True)
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.28); p.paragraph_format.space_after=Pt(7); add_text(p,text)

doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(1); sec.bottom_margin=Inches(1); sec.left_margin=Inches(1); sec.right_margin=Inches(1); sec.header_distance=Inches(.49); sec.footer_distance=Inches(.49)
styles=doc.styles
normal=styles["Normal"]; normal.font.name="Calibri"; normal._element.rPr.rFonts.set(qn("w:ascii"),"Calibri"); normal.font.size=Pt(11); normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.1
for n,sz,col,bef,aft in [("Heading 1",16,PURPLE,16,8),("Heading 2",13,PURPLE,12,6),("Heading 3",12,"3D356E",8,4)]:
    s=styles[n]; s.font.name="Calibri"; s._element.rPr.rFonts.set(qn("w:ascii"),"Calibri"); s.font.size=Pt(sz); s.font.color.rgb=RGBColor.from_string(col); s.font.bold=True; s.paragraph_format.space_before=Pt(bef); s.paragraph_format.space_after=Pt(aft)

# Quiet running furniture
hp=sec.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT; add_text(hp,"MARKOPILOT  |  HOW IT WORKS",size=8,color=MUTED,bold=True)
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; add_text(fp,"Markopilot — Your opportunity engine",size=8,color=MUTED)

# Cover
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(95); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; add_text(p,"A PLAIN-ENGLISH GUIDE",size=10,color=PURPLE,bold=True)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(10); p.paragraph_format.space_after=Pt(10); add_text(p,"How Markopilot\nworks for your business",size=30,color=INK,bold=True)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(26); add_text(p,"How it notices change, connects opportunities, and helps you take the right next step.",size=14,color=MUTED)
callout(doc,"The short version","Markopilot keeps an eye on the market around your business. It turns useful signals and promising leads into clear opportunities, then prepares practical next actions while you remain in control.")
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(46); add_text(p,"MARKOPILOT",size=12,color=PURPLE,bold=True); add_text(p,"  •  Detailed non-technical overview",size=10,color=MUTED)
doc.add_page_break()

heading(doc,"1. What Markopilot is",1)
body(doc,"Markopilot is a growth co-worker for a business. It is built to answer a difficult everyday question: “What should we do next to create more awareness, conversations, and customers?”")
body(doc,"Rather than treating social media, lead lists, market research, and outreach as separate chores, Markopilot brings them into one working loop. It watches for meaningful change, looks for the people and companies connected to that change, and helps your team decide what to do with the opportunity.")
callout(doc,"It is not a generic content machine","A scheduled post is only one possible outcome. Markopilot’s main job is to notice a relevant opening and make that opening easier to act on.")
heading(doc,"What it is designed to help with",2)
for x in ["Keeping track of competitor moves, industry discussions, brand mentions, and market changes.","Finding and scoring prospective customers that match your ideal customer profile.","Giving meaning to the information: why it matters, how urgent it is, and what it may be worth doing.","Preparing practical actions such as a social post, a helpful reply, an alert, or a reviewed outreach step.","Maintaining a steady content rhythm even when there is no major news event."]:
    bullet(doc,x)
heading(doc,"2. The idea behind it: one growth loop",1)
body(doc,"Many businesses have useful information scattered everywhere. A competitor changes its price. A potential customer publicly describes a problem. A lead appears in a search. A post performs unexpectedly well. Each is useful, but on its own it is easy to miss or ignore.")
body(doc,"Markopilot treats these as inputs to one shared process. Its central concept is an opportunity: a situation that could reasonably help the business grow, protect its position, or start a valuable conversation.")
numbered(doc,"1","It notices","Markopilot gathers relevant signals from sources such as news and industry feeds, search, public discussions, competitor pages, and mentions.")
numbered(doc,"2","It connects","It combines the signal with your business context: your industry, target audience, growth goals, locations, and the problems you solve. It also discovers and scores suitable prospects independently of news.")
numbered(doc,"3","It decides","Strong signals and qualified leads are turned into opportunities. Each opportunity records what happened, why it matters, how relevant it is, and how quickly it may need attention.")
numbered(doc,"4","It helps you move","For each opportunity, Markopilot can prepare a response or recommendation. Depending on your settings, that may be queued for review or allowed to move ahead automatically within safe limits.")

heading(doc,"3. What happens in a typical week",1)
body(doc,"The best way to understand Markopilot is to follow a simple example. Imagine a company that sells HR software to growing businesses in East Africa.")
heading(doc,"Monday: a market event appears",2)
body(doc,"Markopilot notices that a competitor has removed an important feature from its lower-priced plan. It does not simply report the news. It asks whether that change creates an opening for the HR software company.")
callout(doc,"Opportunity created","“Reach companies affected by the feature change.” The opportunity is given a relevance score, an urgency level, and an explanation of why it could matter.")
heading(doc,"Tuesday: people and companies are connected to the event",2)
body(doc,"At the same time, the lead discovery engine is looking for businesses that match the company’s ideal customer profile. If it finds high-quality prospects related to the competitor’s market, those leads can be attached to the same opportunity instead of sitting in an isolated list.")
heading(doc,"Wednesday: a recommended response is prepared",2)
body(doc,"Markopilot may recommend a clear positioning post, a helpful reply to a relevant public conversation, or a notification that invites the team to consider a direct approach. It explains the reason for the recommendation rather than leaving the team to interpret raw data.")
heading(doc,"Thursday: the business decides how to proceed",2)
body(doc,"If your settings require review, the prepared action waits for a person. You can approve it, reject it, or decide to do nothing. If you have allowed safe automatic actions, the relevant content can be queued without waiting for a manual handoff.")
heading(doc,"Friday: the loop gets smarter in practice",2)
body(doc,"The outcomes of actions—such as engagement, replies, or interest—can become useful context for future decisions. A topic that generates strong attention may be worth developing further. A market signal that produces no response can be treated as less important next time.")
heading(doc,"4. The different ways Markopilot finds opportunities",1)
heading(doc,"Market signals",2)
body(doc,"These are things happening around your business: a competitor announcement, a customer conversation, an industry trend, a policy change, a news story, or a relevant discussion online. Markopilot filters out the background noise and focuses on signals that have a clear possible connection to your business goals.")
heading(doc,"Qualified leads",2)
body(doc,"Lead discovery works from your ideal customer profile. It searches for people and companies that may fit, extracts the useful details, and gives each lead a score. Strong leads are not treated as a pile of names: they can become opportunities with a clear reason for attention.")
heading(doc,"Your own brand context",2)
body(doc,"Your description, audience, positioning, content pillars, target locations, and growth goal act as Markopilot’s compass. They help it decide whether a signal is truly useful for your business instead of merely popular or interesting in general.")

heading(doc,"5. What Markopilot can do once it sees an opportunity",1)
body(doc,"Markopilot does not assume that every opportunity should become a public post. It can suggest different next steps depending on the situation and your preferences.")
for title,text in [("Prepare a timely post","Turn a market change, product insight, or useful viewpoint into an on-brand post for the social channels you have connected."),("Prepare a helpful reply","Draft a thoughtful response to a relevant public conversation without asking your team to start from a blank page."),("Surface a high-value prospect","Bring a qualified company or person to your attention together with the reason they are a good fit."),("Create an alert","Tell the business owner about an event that may need a human decision, such as a sensitive competitor move or important market change."),("Keep evergreen content moving","Create regular content from your chosen themes even when no urgent market event is happening.")]:
    heading(doc,title,3); body(doc,text)
callout(doc,"A useful distinction","Markopilot prepares and organizes actions. The level of automatic execution is always chosen by you. The system should feel like a capable co-worker, not an uncontrolled black box.")
heading(doc,"6. You remain in control",1)
body(doc,"Every business has a different comfort level. Some teams want to review every draft. Others want routine social content to move quickly while keeping outreach under review. Markopilot gives you a simple autonomy choice so the system fits your operating style.")
for x in ["Review-first: drafts and recommendations wait for your approval.","Supervised: safe content can be queued while more sensitive actions remain for review.","More autonomous: Markopilot can move approved types of work forward within the boundaries you set."]:
    bullet(doc,x)
body(doc,"You can change these settings as your trust in the system grows. The point is not to remove people from the process; it is to remove the delay and busywork that stop good opportunities from being noticed.")
heading(doc,"7. What Markopilot does not promise",1)
body(doc,"Markopilot is designed to improve focus and speed, not to guarantee sales or replace sound business judgment. Not every signal deserves a response, and not every lead will become a customer. The value is in helping your team see the strongest openings sooner, with better context and less manual research.")
body(doc,"It should also be used responsibly. Outreach and public communication still benefit from human judgment, especially when the message is sensitive, the audience is regulated, or the business is protecting a carefully built reputation.")

heading(doc,"8. Getting the most from Markopilot",1)
body(doc,"Markopilot becomes more useful when it understands what good growth means for your business. A few thoughtful setup choices make a major difference.")
for n,t,x in [("1","Be specific about the goal","Describe what you are trying to achieve: more qualified conversations, a new market position, better awareness in a sector, or a particular type of customer."),("2","Describe your best-fit customer","Include the industries, roles, company types, locations, and common problems that make someone a strong fit."),("3","Set a clear voice","Give the system the tone you want it to use: direct, professional, playful, reassuring, technical, or something else."),("4","Start with review","For a new brand, it is sensible to review the first recommendations. This helps you refine the boundaries and gives Markopilot a clearer working rhythm."),("5","Use the opportunity view regularly","The value is not just in the actions that occur. It is also in seeing what the market is telling you, where attention is building, and which ideas are worth discussing.")]: numbered(doc,n,t,x)
heading(doc,"9. A simple way to think about Markopilot",1)
callout(doc,"The one-sentence description","Markopilot is an autonomous opportunity discovery and action engine that helps a business notice what matters, understand why it matters, and take a well-timed next step.")
body(doc,"Social publishing is one action. Lead discovery is another. Market research, replies, alerts, and follow-up are others. What makes Markopilot different is the layer above them: it tries to connect these activities to a shared business objective, so the work is coordinated rather than scattered.")
heading(doc,"Glossary",1)
for term,definition in [("Signal","A piece of information that may be relevant to your business, such as a competitor change or market discussion."),("Lead","A person or company that may match your ideal customer profile."),("Opportunity","A situation that could advance your business goals and deserves consideration."),("Action","A practical next step Markopilot prepares or recommends, such as a draft post, a reply, or an alert."),("Autonomy level","The amount of permission you give Markopilot to move work forward without waiting for approval.")]:
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(4); add_text(p,term+": ",bold=True,color=PURPLE); add_text(p,definition)
body(doc,"This guide describes Markopilot in plain language. The exact actions available to a brand depend on the channels it connects and the settings it chooses.")

OUT.parent.mkdir(exist_ok=True); doc.core_properties.title="How Markopilot Works"; doc.core_properties.subject="A detailed non-technical guide"; doc.save(OUT)
print(OUT)
